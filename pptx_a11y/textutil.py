from pptx.enum.dml import MSO_FILL


def fill_rgb(fill_owner):
    """Return the solid-fill RGB (r,g,b) of a shape or background object, or
    None if it has no resolvable solid RGB fill (no fill, gradient, picture,
    or a theme/inherited color we cannot resolve to concrete RGB)."""
    try:
        fill = fill_owner.fill
        if fill.type == MSO_FILL.SOLID:
            rgb = fill.fore_color.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:  # noqa: BLE001 - theme/absent fills are not resolvable here
        return None
    return None


def iter_runs(prs):
    # Descend into groups (via iter_shapes) so text nested inside grouped
    # shapes is checked too, not just top-level shapes.
    from pptx_a11y.checks import iter_shapes
    for i, shape in iter_shapes(prs):
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                yield i, shape, para, run


def run_rgb(run):
    try:
        color = run.font.color
        if color and color.type is not None and color.rgb is not None:
            rgb = color.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:  # noqa: BLE001 - theme/inherited colors are unresolvable here
        return None
    return None


def run_pt(run) -> float | None:
    sz = run.font.size
    return sz.pt if sz is not None else None
